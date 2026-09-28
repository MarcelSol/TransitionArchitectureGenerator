"""
excel_reader.py

Reads a machine-readable TAG Excel workbook and reconstructs
a TransitionModel.

The Summary and Validation sheets are intentionally ignored.
"""

from pathlib import Path

from openpyxl import load_workbook

from tag.transition_model import (
    TransitionModel,
    TransitionNode,
    TransitionChild,
    TransitionInterface,
    NodeCategory,
    InterfaceDirection,
    TransferType,
    InterfaceKey,
)


class ExcelReader:

    REQUIRED_SHEETS = {
        "Milestones",
        "Nodes",
        "Interfaces",
        "Children",
    }

    @staticmethod
    def read(
        filename: str,
    ) -> TransitionModel:

        path = Path(filename)

        if not path.exists():
            raise FileNotFoundError(
                f"Excel file not found: {filename}"
            )

        if path.suffix.lower() != ".xlsx":
            raise ValueError(
                f"Expected an .xlsx file: {filename}"
            )

        workbook = load_workbook(
            filename,
            data_only=True,
        )

        ExcelReader._validate_sheets(
            workbook.sheetnames
        )

        model = TransitionModel()

        ExcelReader._read_milestones(
            workbook,
            model,
        )

        ExcelReader._read_nodes(
            workbook,
            model,
        )

        ExcelReader._read_children(
            workbook,
            model,
        )

        ExcelReader._read_interfaces(
            workbook,
            model,
        )

        return model

    # -----------------------------------------------------------------
    # Workbook validation
    # -----------------------------------------------------------------

    @staticmethod
    def _validate_sheets(
        sheet_names: list[str],
    ) -> None:

        available = set(sheet_names)

        missing = (
            ExcelReader.REQUIRED_SHEETS
            - available
        )

        if missing:

            missing_text = ", ".join(
                sorted(missing)
            )

            raise ValueError(
                "Excel workbook is missing "
                f"required sheet(s): {missing_text}"
            )

    # -----------------------------------------------------------------
    # Milestones
    # -----------------------------------------------------------------

    @staticmethod
    def _read_milestones(
        workbook,
        model: TransitionModel,
    ) -> None:

        sheet = workbook["Milestones"]

        ExcelReader._require_headers(
            sheet,
            [
                "Order",
                "Milestone",
            ],
        )

        rows = []

        for row_number in range(
            2,
            sheet.max_row + 1,
        ):

            order = sheet.cell(
                row_number,
                1,
            ).value

            milestone = sheet.cell(
                row_number,
                2,
            ).value

            if (
                order is None
                and milestone is None
            ):
                continue

            if order is None:
                raise ValueError(
                    "Milestones sheet: "
                    f"row {row_number} has no Order."
                )

            if milestone is None:
                raise ValueError(
                    "Milestones sheet: "
                    f"row {row_number} has no Milestone."
                )

            rows.append(
                (
                    int(order),
                    str(milestone),
                )
            )

        rows.sort(
            key=lambda item: item[0]
        )

        expected_order = 1

        for order, milestone in rows:

            if order != expected_order:
                raise ValueError(
                    "Milestones sheet: expected "
                    f"Order {expected_order}, "
                    f"found {order}."
                )

            if milestone in model.milestones:
                raise ValueError(
                    "Milestones sheet: duplicate "
                    f"milestone '{milestone}'."
                )

            model.milestones.append(
                milestone
            )

            expected_order += 1

        if not model.milestones:
            raise ValueError(
                "Milestones sheet contains no milestones."
            )

    # -----------------------------------------------------------------
    # Nodes
    # -----------------------------------------------------------------

    @staticmethod
    def _read_nodes(
        workbook,
        model: TransitionModel,
    ) -> None:

        sheet = workbook["Nodes"]

        ExcelReader._require_headers(
            sheet,
            [
                "ID",
                "Name",
                "Category",
                "Width",
                "Height",
                "First Appears",
                "Retired In",
                "Visible On",
            ],
        )

        for row_number in range(
            2,
            sheet.max_row + 1,
        ):

            if ExcelReader._row_is_empty(
                sheet,
                row_number,
            ):
                continue

            node_id = ExcelReader._required_string(
                sheet,
                row_number,
                "ID",
            )

            name = ExcelReader._required_string(
                sheet,
                row_number,
                "Name",
            )

            category_text = ExcelReader._required_string(
                sheet,
                row_number,
                "Category",
            )

            category = (
                ExcelReader._node_category(
                    category_text,
                    row_number,
                )
            )

            width = ExcelReader._required_number(
                sheet,
                row_number,
                "Width",
            )

            height = ExcelReader._required_number(
                sheet,
                row_number,
                "Height",
            )

            first_appears = (
                ExcelReader._optional_string(
                    sheet,
                    row_number,
                    "First Appears",
                )
            )

            retired_in = (
                ExcelReader._optional_string(
                    sheet,
                    row_number,
                    "Retired In",
                )
            )

            visible_on = (
                ExcelReader._read_visibility(
                    sheet,
                    row_number,
                    model.milestones,
                    "Visible On",
                )
            )

            if node_id in model.nodes:
                raise ValueError(
                    "Nodes sheet: duplicate "
                    f"node ID '{node_id}' "
                    f"at row {row_number}."
                )

            ExcelReader._validate_milestone_value(
                first_appears,
                model.milestones,
                "First Appears",
                row_number,
            )

            ExcelReader._validate_milestone_value(
                retired_in,
                model.milestones,
                "Retired In",
                row_number,
            )

            model.nodes[node_id] = (
                TransitionNode(
                    id=node_id,
                    name=name,
                    category=category,
                    width=width,
                    height=height,
                    visible_on=visible_on,
                    first_appears=first_appears,
                    retired_in=retired_in,
                )
            )

    # -----------------------------------------------------------------
    # Children
    # -----------------------------------------------------------------

    @staticmethod
    def _read_children(
        workbook,
        model: TransitionModel,
    ) -> None:

        sheet = workbook["Children"]

        ExcelReader._require_headers(
            sheet,
            [
                "Parent ID",
                "ID",
                "Name",
                "Category",
                "X",
                "Y",
                "Width",
                "Height",
                "Visible On",
            ],
        )

        child_ids = set()

        for row_number in range(
            2,
            sheet.max_row + 1,
        ):

            if ExcelReader._row_is_empty(
                sheet,
                row_number,
            ):
                continue

            parent_id = ExcelReader._required_string(
                sheet,
                row_number,
                "Parent ID",
            )

            child_id = ExcelReader._required_string(
                sheet,
                row_number,
                "ID",
            )

            name = ExcelReader._required_string(
                sheet,
                row_number,
                "Name",
            )

            category_text = ExcelReader._required_string(
                sheet,
                row_number,
                "Category",
            )

            category = (
                ExcelReader._node_category(
                    category_text,
                    row_number,
                )
            )

            x = ExcelReader._required_number(
                sheet,
                row_number,
                "X",
            )

            y = ExcelReader._required_number(
                sheet,
                row_number,
                "Y",
            )

            width = ExcelReader._required_number(
                sheet,
                row_number,
                "Width",
            )

            height = ExcelReader._required_number(
                sheet,
                row_number,
                "Height",
            )

            visible_on = (
                ExcelReader._read_visibility(
                    sheet,
                    row_number,
                    model.milestones,
                    "Visible On",
                )
            )

            if parent_id not in model.nodes:
                raise ValueError(
                    "Children sheet: parent node "
                    f"'{parent_id}' does not exist "
                    f"at row {row_number}."
                )

            if child_id in child_ids:
                raise ValueError(
                    "Children sheet: duplicate "
                    f"child ID '{child_id}' "
                    f"at row {row_number}."
                )

            if child_id in model.nodes:
                raise ValueError(
                    "Children sheet: child ID "
                    f"'{child_id}' is also a "
                    "top-level node ID."
                )

            child_ids.add(child_id)

            child = TransitionChild(
                id=child_id,
                name=name,
                category=category,
                x=x,
                y=y,
                width=width,
                height=height,
                visible_on=visible_on,
            )

            model.nodes[
                parent_id
            ].children.append(
                child
            )

    # -----------------------------------------------------------------
    # Interfaces
    # -----------------------------------------------------------------

    @staticmethod
    def _read_interfaces(
        workbook,
        model: TransitionModel,
    ) -> None:

        sheet = workbook["Interfaces"]

        ExcelReader._require_headers(
            sheet,
            [
                "ID",
                "Label",
                "Source",
                "Target",
                "Direction",
                "Transfer Type",
                "First Appears",
                "Retired In",
                "Visible On",
            ],
        )

        for row_number in range(
            2,
            sheet.max_row + 1,
        ):

            if ExcelReader._row_is_empty(
                sheet,
                row_number,
            ):
                continue

            interface_id = (
                ExcelReader._required_string(
                    sheet,
                    row_number,
                    "ID",
                )
            )

            label = (
                ExcelReader._optional_string(
                    sheet,
                    row_number,
                    "Label",
                )
            )

            source = (
                ExcelReader._required_string(
                    sheet,
                    row_number,
                    "Source",
                )
            )

            target = (
                ExcelReader._required_string(
                    sheet,
                    row_number,
                    "Target",
                )
            )

            direction_text = (
                ExcelReader._required_string(
                    sheet,
                    row_number,
                    "Direction",
                )
            )

            transfer_text = (
                ExcelReader._required_string(
                    sheet,
                    row_number,
                    "Transfer Type",
                )
            )

            first_appears = (
                ExcelReader._optional_string(
                    sheet,
                    row_number,
                    "First Appears",
                )
            )

            retired_in = (
                ExcelReader._optional_string(
                    sheet,
                    row_number,
                    "Retired In",
                )
            )

            visible_on = (
                ExcelReader._read_visibility(
                    sheet,
                    row_number,
                    model.milestones,
                    "Visible On",
                )
            )

            if source not in model.nodes:
                raise ValueError(
                    "Interfaces sheet: source "
                    f"node '{source}' does not exist "
                    f"at row {row_number}."
                )

            if target not in model.nodes:
                raise ValueError(
                    "Interfaces sheet: target "
                    f"node '{target}' does not exist "
                    f"at row {row_number}."
                )

            direction = (
                ExcelReader._interface_direction(
                    direction_text,
                    row_number,
                )
            )

            transfer_type = (
                ExcelReader._transfer_type(
                    transfer_text,
                    row_number,
                )
            )

            ExcelReader._validate_milestone_value(
                first_appears,
                model.milestones,
                "First Appears",
                row_number,
            )

            ExcelReader._validate_milestone_value(
                retired_in,
                model.milestones,
                "Retired In",
                row_number,
            )

            key = InterfaceKey(
                source=source,
                target=target,
                direction=direction,
                transfer_type=transfer_type,
            )

            if key in model.interfaces:
                raise ValueError(
                    "Interfaces sheet: duplicate "
                    "interface definition for "
                    f"'{source}' -> '{target}' "
                    f"({direction_text}, "
                    f"{transfer_text}) "
                    f"at row {row_number}."
                )

            interface = TransitionInterface(
                id=interface_id,
                source=source,
                target=target,
                direction=direction,
                transfer_type=transfer_type,
                visible_on=visible_on,
                first_appears=first_appears,
                retired_in=retired_in,
                label=label or None,
            )

            model.interfaces[key] = interface

    # -----------------------------------------------------------------
    # Helpers
    # -----------------------------------------------------------------

    @staticmethod
    def _require_headers(
        sheet,
        required_headers: list[str],
    ) -> None:

        headers = {
            sheet.cell(
                1,
                column,
            ).value
            for column in range(
                1,
                sheet.max_column + 1,
            )
        }

        missing = [
            header
            for header in required_headers
            if header not in headers
        ]

        if missing:
            raise ValueError(
                f"Sheet '{sheet.title}' is missing "
                "required column(s): "
                + ", ".join(missing)
            )

    # -----------------------------------------------------------------

    @staticmethod
    def _column_index(
        sheet,
        header: str,
    ) -> int:

        for column in range(
            1,
            sheet.max_column + 1,
        ):

            if (
                sheet.cell(
                    1,
                    column,
                ).value
                == header
            ):
                return column

        raise ValueError(
            f"Sheet '{sheet.title}' has no "
            f"column '{header}'."
        )

    # -----------------------------------------------------------------

    @staticmethod
    def _cell(
        sheet,
        row_number: int,
        header: str,
    ):

        column = ExcelReader._column_index(
            sheet,
            header,
        )

        return sheet.cell(
            row_number,
            column,
        ).value

    # -----------------------------------------------------------------

    @staticmethod
    def _required_string(
        sheet,
        row_number: int,
        header: str,
    ) -> str:

        value = ExcelReader._cell(
            sheet,
            row_number,
            header,
        )

        if value is None:
            raise ValueError(
                f"Sheet '{sheet.title}': "
                f"row {row_number} has no "
                f"value for '{header}'."
            )

        text = str(value).strip()

        if not text:
            raise ValueError(
                f"Sheet '{sheet.title}': "
                f"row {row_number} has an empty "
                f"value for '{header}'."
            )

        return text

    # -----------------------------------------------------------------

    @staticmethod
    def _optional_string(
        sheet,
        row_number: int,
        header: str,
    ) -> str:

        value = ExcelReader._cell(
            sheet,
            row_number,
            header,
        )

        if value is None:
            return ""

        return str(value).strip()

    # -----------------------------------------------------------------

    @staticmethod
    def _required_number(
        sheet,
        row_number: int,
        header: str,
    ) -> float:

        value = ExcelReader._cell(
            sheet,
            row_number,
            header,
        )

        if value is None:
            raise ValueError(
                f"Sheet '{sheet.title}': "
                f"row {row_number} has no "
                f"value for '{header}'."
            )

        try:
            return float(value)
        except (
            TypeError,
            ValueError,
        ) as exc:

            raise ValueError(
                f"Sheet '{sheet.title}': "
                f"row {row_number} has invalid "
                f"number '{value}' for "
                f"'{header}'."
            ) from exc

    # -----------------------------------------------------------------

    @staticmethod
    def _node_category(
        value: str,
        row_number: int,
    ) -> NodeCategory:

        for category in NodeCategory:

            if category.value == value:
                return category

        raise ValueError(
            "Invalid node category "
            f"'{value}' at row {row_number}."
        )

    # -----------------------------------------------------------------

    @staticmethod
    def _interface_direction(
        value: str,
        row_number: int,
    ) -> InterfaceDirection:

        normalized = value.strip().lower()

        if normalized in {
            "one-way",
            "one way",
            "one_way",
        }:
            return InterfaceDirection.ONE_WAY

        if normalized in {
            "two-way",
            "two way",
            "two_way",
        }:
            return InterfaceDirection.TWO_WAY

        raise ValueError(
            "Invalid interface direction "
            f"'{value}' at row {row_number}."
        )

    # -----------------------------------------------------------------

    @staticmethod
    def _transfer_type(
        value: str,
        row_number: int,
    ) -> TransferType:

        normalized = value.strip().lower()

        if normalized in {
            "automatic",
            "automated",
        }:
            return TransferType.AUTOMATED

        if normalized == "manual":
            return TransferType.MANUAL

        raise ValueError(
            "Invalid transfer type "
            f"'{value}' at row {row_number}."
        )

    # -----------------------------------------------------------------

    @staticmethod
    def _read_visibility(
        sheet,
        row_number: int,
        milestones: list[str],
        header: str,
    ) -> set[str]:

        value = ExcelReader._cell(
            sheet,
            row_number,
            header,
        )

        if value is None:
            return set()

        text = str(value).strip()

        if not text:
            return set()

        values = {
            item.strip()
            for item in text.split(";")
            if item.strip()
        }

        unknown = (
            values
            - set(milestones)
        )

        if unknown:

            unknown_text = ", ".join(
                sorted(unknown)
            )

            raise ValueError(
                f"Sheet '{sheet.title}': "
                f"row {row_number} contains "
                f"unknown milestone(s) in "
                f"'{header}': {unknown_text}"
            )

        return values

    # -----------------------------------------------------------------

    @staticmethod
    def _validate_milestone_value(
        value: str,
        milestones: list[str],
        field_name: str,
        row_number: int,
    ) -> None:

        if not value:
            return

        if value not in milestones:
            raise ValueError(
                f"Invalid milestone '{value}' "
                f"in '{field_name}' at row "
                f"{row_number}."
            )

    # -----------------------------------------------------------------

    @staticmethod
    def _row_is_empty(
        sheet,
        row_number: int,
    ) -> bool:

        for column in range(
            1,
            sheet.max_column + 1,
        ):

            value = sheet.cell(
                row_number,
                column,
            ).value

            if value is not None:
                return False

        return True
